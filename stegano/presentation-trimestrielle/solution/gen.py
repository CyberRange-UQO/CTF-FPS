#!/usr/bin/env python3
"""Genere le defi presentation-trimestrielle.

Fuite reelle et frequente : une image est "rognee" dans PowerPoint pour masquer
une zone sensible. Le rognage n'est PAS destructif -> l'image complete reste
dans ppt/media/. On la recupere en dezippant le .pptx (ou via "Reinitialiser
l'image"). Le bas de l'image (masque par le rognage) contient le flag.
"""
import os

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches

HERE = os.path.dirname(os.path.abspath(__file__))
DEFI = os.path.normpath(os.path.join(HERE, "..", "defi"))
FLAG = "FPS{rognage_powerpoint_non_destructif_image_intacte}"

FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FONTB = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"

W, H = 1100, 760
CROP_BOTTOM = 0.32  # fraction basse masquee par le rognage


def build_image(path):
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    big = ImageFont.truetype(FONTB, 34)
    med = ImageFont.truetype(FONTB, 24)
    small = ImageFont.truetype(FONT, 20)

    d.rectangle([0, 0, W, 70], fill=(31, 73, 125))
    d.text((24, 18), "Resultats trimestriels - T3 2026", font=big, fill="white")

    # Faux histogramme des ventes par region.
    regions = [("Outaouais", 180), ("Montreal", 240), ("Quebec", 150), ("Est", 90)]
    x0, base, bw, gap = 90, 430, 120, 90
    maxv = 260
    for i, (name, val) in enumerate(regions):
        x = x0 + i * (bw + gap)
        top = base - int((val / maxv) * 300)
        d.rectangle([x, top, x + bw, base], fill=(79, 129, 189))
        d.text((x + 10, base + 10), name, font=small, fill="black")
        d.text((x + 24, top - 28), str(val), font=med, fill="black")
    d.line([70, base, W - 60, base], fill="black", width=2)
    d.line([70, 120, 70, base], fill="black", width=2)

    # Bande basse sensible (sera masquee par le rognage dans PowerPoint).
    d.rectangle([0, int(H * (1 - CROP_BOTTOM)), W, H], fill=(245, 230, 230))
    yy = int(H * (1 - CROP_BOTTOM)) + 18
    d.text((24, yy), "NOTE CONFIDENTIELLE - ne pas diffuser a l'externe",
           font=med, fill=(150, 0, 0))
    d.text((24, yy + 40),
           "Marge nette reelle T3 : 4,2 % (chiffre non communique au public).",
           font=small, fill="black")
    d.text((24, yy + 74), "Reference interne : " + FLAG, font=small, fill="black")

    img.save(path)


def main():
    full = os.path.join(HERE, "image_complete.png")
    build_image(full)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # vierge

    pic = slide.shapes.add_picture(full, Inches(1.2), Inches(0.9), width=Inches(10.9))
    # Rognage non destructif : on masque le bas (ou se trouve la note).
    pic.crop_bottom = CROP_BOTTOM

    # Un peu de contexte texte autour.
    tx = slide.shapes.add_textbox(Inches(1.2), Inches(6.7), Inches(10), Inches(0.6))
    tx.text_frame.text = "Presentation interne - Direction commerciale - T3 2026"

    out = os.path.join(DEFI, "presentation-T3-2026.pptx")
    prs.save(out)
    os.remove(full)
    print("genere:", out)


if __name__ == "__main__":
    main()
